from docx import Document
import os
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH


document = Document()
items = []
for i, j, k in os.walk('./Results 4/'):
    for i in k:
        item = list()
        h = i.split(' ')
        item.extend([h[0], h[1], h[4], h[2], h[3]])
        with open(f'./Results 4/{i}', 'r') as f:
            data = f.read()
            data = data.split('\n')
            item.append(data[0].split(' ')[1])
            item.append(data[1].split(' ')[1])
            item.append(data[2].split(' ')[1])
            item.append(data[3].split(' ')[1])
        items.append(item)
print(items)
items.sort(key=lambda x: x[8])
table = document.add_table(1, 9)
for i in range(9):
    table.cell(0, i).vertical_alignment = WD_ALIGN_VERTICAL.CENTER

heading_cells = table.rows[0].cells
heading_cells[0].text = 'Crypto-Currency'
heading_cells[1].text = 'Network'
heading_cells[2].text = 'Optimizer'
heading_cells[3].text = 'BS'
heading_cells[4].text = 'TS'
heading_cells[5].text = 'MSE'
heading_cells[6].text = 'RMSE'
heading_cells[7].text = 'MAE'
heading_cells[8].text = 'MSLE'


for index, item in enumerate(items):
    row = table.add_row()
    cells = row.cells
    for i in range(9):
        table.cell(index, i).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cells[0].text = str(item[0])
    cells[1].text = str(item[1])
    cells[2].text = str(item[2])
    cells[3].text = str(item[3])
    cells[4].text = str(item[4])
    cells[5].text = str(round(float(item[5]), 5))
    cells[6].text = str(round(float(item[6]), 5))
    cells[7].text = str(round(float(item[7]), 5))
    cells[8].text = str(round(float(item[8]), 5))


# table.style = 'ColorfulGrid-Accent1'
table.style = 'DarkList-Accent1'
# table.style = 'ColorfulShading-Accent1'

document.save('Results MSLE.docx')
